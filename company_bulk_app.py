import os
import csv
import json
from pathlib import Path
from io import BytesIO, StringIO
from datetime import datetime

from flask import Flask, request, jsonify, send_file
from pypdf import PdfReader
from openai import OpenAI, AuthenticationError, BadRequestError, APIStatusError
from dotenv import load_dotenv

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(dotenv_path=BASE_DIR / ".env", override=True)

app = Flask(__name__)

last_bulk_analysis = {
    "timestamp": None,
    "job_description": "",
    "results": [],
}


def _clean_env_key(value: str) -> str:
    return (value or "").strip().strip('"').strip("'")


def _is_real_api_key(value: str) -> bool:
    if not value:
        return False
    lowered = value.lower()
    placeholders = (
        "your-key-here",
        "your-openai-key-here",
        "your-perplexity-key-here",
        "replace-me",
        "changeme",
        "dummy",
        "test",
    )
    return not any(token in lowered for token in placeholders)


pplx_api_key = _clean_env_key(os.getenv("PPLX_API_KEY"))
openai_api_key = _clean_env_key(os.getenv("OPENAI_API_KEY"))


def _build_clients():
    available = []
    if _is_real_api_key(pplx_api_key):
        available.append(("perplexity", OpenAI(api_key=pplx_api_key, base_url="https://api.perplexity.ai")))
    if _is_real_api_key(openai_api_key):
        available.append(("openai", OpenAI(api_key=openai_api_key)))
    return available


clients = _build_clients()


def extract_resume_text(file_storage):
    filename = (file_storage.filename or "").lower()

    if filename.endswith(".pdf"):
        reader = PdfReader(file_storage)
        text = []
        for page in reader.pages:
            text.append(page.extract_text() or "")
        return "\n".join(text).strip()

    if filename.endswith(".txt"):
        return file_storage.read().decode("utf-8", errors="ignore").strip()

    if filename.endswith(".docx") and DOCX_AVAILABLE:
        file_storage.stream.seek(0)
        doc = Document(file_storage.stream)
        return "\n".join([p.text for p in doc.paragraphs]).strip()

    return ""


def _safe_json_parse(text: str):
    try:
        return json.loads(text)
    except Exception:
        start = text.find("{")
        end = text.rfind("}")
        if start != -1 and end != -1 and end > start:
            try:
                return json.loads(text[start:end + 1])
            except Exception:
                return None
        return None


def analyze_resume(resume_text: str, job_description: str):
    prompt = f"""
Compare this resume against the job description.
Return ONLY valid JSON with keys:
- match_score (integer 0-100)
- strengths (array of 3 short bullet strings)
- missing_keywords (array of up to 8 strings)
- improvement_suggestions (array of 3 short bullet strings)
- summary (2-3 lines)

Resume:
{resume_text}

Job Description:
{job_description}
"""

    auth_failures = 0
    last_error = None

    for active_provider, active_client in clients:
        try:
            if active_provider == "perplexity":
                response = active_client.chat.completions.create(
                    model="sonar-pro",
                    messages=[
                        {"role": "system", "content": "You are an ATS and recruiting assistant."},
                        {"role": "user", "content": prompt},
                    ],
                )
                output_text = (response.choices[0].message.content or "").strip()
            else:
                response = active_client.responses.create(
                    model="gpt-4.1-mini",
                    input=prompt,
                )
                output_text = (response.output_text or "").strip()

            parsed = _safe_json_parse(output_text)
            if parsed:
                score = parsed.get("match_score", 0)
                try:
                    score = int(score)
                except Exception:
                    score = 0

                return {
                    "match_score": max(0, min(score, 100)),
                    "strengths": parsed.get("strengths", []) if isinstance(parsed.get("strengths", []), list) else [],
                    "missing_keywords": parsed.get("missing_keywords", []) if isinstance(parsed.get("missing_keywords", []), list) else [],
                    "improvement_suggestions": parsed.get("improvement_suggestions", []) if isinstance(parsed.get("improvement_suggestions", []), list) else [],
                    "summary": str(parsed.get("summary", "")).strip(),
                    "raw_analysis": output_text,
                }

            return {
                "match_score": 0,
                "strengths": [],
                "missing_keywords": [],
                "improvement_suggestions": [],
                "summary": output_text[:500],
                "raw_analysis": output_text,
            }
        except AuthenticationError:
            auth_failures += 1
            continue
        except Exception as exc:
            last_error = str(exc)
            continue

    if auth_failures == len(clients) and len(clients) > 0:
        raise AuthenticationError("Invalid API key for all configured providers.")

    if last_error:
        raise RuntimeError(last_error)

    raise RuntimeError("No AI provider returned a response.")


@app.route("/", methods=["GET"])
def home():
    html_path = BASE_DIR / "company_index.html"
    if html_path.exists():
        with open(html_path, "r", encoding="utf-8") as file:
            return file.read()
    return "<h1>Company Bulk Resume Scanner</h1><p>company_index.html not found.</p>", 404


@app.route("/favicon.ico", methods=["GET"])
def favicon():
    return "", 204


@app.route("/bulk-scan", methods=["POST"])
def bulk_scan():
    try:
        if not clients:
            return jsonify({
                "error": "API key not configured. Add valid PPLX_API_KEY or OPENAI_API_KEY in .env and restart this app."
            }), 500

        resumes = request.files.getlist("resumes")
        job_description = (request.form.get("job_description") or "").strip()

        if not job_description:
            return jsonify({"error": "Job description is required."}), 400

        if not resumes:
            return jsonify({"error": "Please upload resume files."}), 400

        total_files = len(resumes)
        if total_files < 1 or total_files > 1000:
            return jsonify({"error": "Company bulk mode supports 1 to 1000 resumes per run."}), 400

        results = []
        failed = []

        for index, resume_file in enumerate(resumes, start=1):
            file_name = resume_file.filename or f"resume_{index}"
            try:
                resume_text = extract_resume_text(resume_file)
                if not resume_text:
                    failed.append({"file_name": file_name, "error": "Could not extract text (supported: PDF/TXT/DOCX)."})
                    continue

                analysis = analyze_resume(resume_text, job_description)
                results.append({
                    "file_name": file_name,
                    "match_score": analysis.get("match_score", 0),
                    "summary": analysis.get("summary", ""),
                    "strengths": analysis.get("strengths", []),
                    "missing_keywords": analysis.get("missing_keywords", []),
                    "improvement_suggestions": analysis.get("improvement_suggestions", []),
                    "status": "processed",
                })
            except Exception as exc:
                failed.append({"file_name": file_name, "error": str(exc)})

        results_sorted = sorted(results, key=lambda item: item.get("match_score", 0), reverse=True)

        global last_bulk_analysis
        last_bulk_analysis = {
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "job_description": job_description,
            "results": results_sorted,
        }

        return jsonify({
            "timestamp": last_bulk_analysis["timestamp"],
            "total_uploaded": total_files,
            "processed": len(results_sorted),
            "failed": len(failed),
            "top_candidates": results_sorted[:10],
            "results": results_sorted,
            "failures": failed,
        })

    except AuthenticationError:
        return jsonify({"error": "Invalid API key (401). Update .env and restart this app."}), 401
    except BadRequestError as exc:
        return jsonify({"error": f"OpenAI request error: {exc}"}), 400
    except APIStatusError as exc:
        return jsonify({"error": f"OpenAI API status error ({exc.status_code}): {exc}"}), 502
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.route("/bulk-export-csv", methods=["GET"])
def bulk_export_csv():
    if not last_bulk_analysis.get("results"):
        return jsonify({"error": "No bulk analysis available. Run a bulk scan first."}), 400

    try:
        csv_buffer = StringIO()
        writer = csv.writer(csv_buffer)

        writer.writerow(["Generated", last_bulk_analysis.get("timestamp", "N/A")])
        writer.writerow([])
        writer.writerow([
            "Rank",
            "File Name",
            "Match Score",
            "Summary",
            "Strengths",
            "Missing Keywords",
            "Improvement Suggestions",
        ])

        for rank, row in enumerate(last_bulk_analysis.get("results", []), start=1):
            writer.writerow([
                rank,
                row.get("file_name", ""),
                row.get("match_score", 0),
                row.get("summary", ""),
                " | ".join(row.get("strengths", [])),
                " | ".join(row.get("missing_keywords", [])),
                " | ".join(row.get("improvement_suggestions", [])),
            ])

        csv_bytes = BytesIO(csv_buffer.getvalue().encode("utf-8-sig"))
        csv_bytes.seek(0)

        stamp = (last_bulk_analysis.get("timestamp") or "report").replace(" ", "_").replace(":", "-")
        filename = f"bulk_resume_analysis_{stamp}.csv"

        return send_file(
            csv_bytes,
            mimetype="text/csv",
            as_attachment=True,
            download_name=filename,
        )
    except Exception as exc:
        return jsonify({"error": f"CSV export failed: {exc}"}), 500


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5001, debug=True)
